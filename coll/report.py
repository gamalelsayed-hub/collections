"""Monthly summary of a cycle (dashboard) and the month-end Excel report."""
from . import stages

LETTER_GROUPS = [("1003", ("1003",)), ("1004", ("1004", "1004PN")), ("1005/1006", ("1005", "1006"))]


def matched_label(v):
    return {1: "مطابق", 0: "غير مطابق"}.get(v, "لم يرجع")


def customer_rows(db, cycle):
    """One row per customer required in the cycle, with the whole month's receipts and its letters."""
    cid = cycle["id"]
    paid = db.paid_by_code(cid)                      # all receipts of the cycle (not only the screen's period)
    letters = {}
    for l in db.letters(cycle_id=cid):
        letters.setdefault(l["code"], []).append(l)
    sent = {r[0] for r in db.con.execute("SELECT DISTINCT code FROM events WHERE cycle_id=? AND kind='إرسال كشف'",
                                         (cid,))}
    neg = stages.negotiating_codes(db)
    thr = stages.thresholds_by_level(db)
    fin = stages.threshold(db, "finance")
    reps = {}
    out = []
    for c in db.cycle_customers(cid, required_only=True):
        if c["rep_name"] not in reps:
            rep = db.rep_for(c["rep_name"])
            reps[c["rep_name"]] = rep["name"] if rep else (c["rep_name"] or "بدون مندوب")
        p = paid.get(c["code"], 0.0)
        base = c["overdue"] or 0
        pct = p / base if base > 0.005 else (1.0 if p else 0.0)
        ls = sorted(letters.get(c["code"], []), key=lambda l: l["id"])
        active = [l for l in ls if l["status"] != "ملغي"]
        level = c["level"] or 0
        out.append({
            **c, "rep": reps[c["rep_name"]], "paid_month": p, "pct_month": pct,
            "matched_text": matched_label(c["matched"]), "sent": c["code"] in sent,
            "letters": ls, "active_letters": active,
            "last_letter": active[-1] if active else None,
            "exited": pct >= fin and not level,
            "stop_sale": level >= 2 and pct < thr[level] and not c["excluded"] and c["code"] not in neg,
            "negotiating": c["code"] in neg,
        })
    return out


def summarize(rows):
    """Totals for a list of customer rows (the whole cycle or one rep)."""
    n = len(rows)
    overdue = sum(r["overdue"] or 0 for r in rows)
    paid = sum(r["paid_month"] for r in rows)
    s = {
        "customers": n, "balance": sum(r["balance"] or 0 for r in rows), "overdue": overdue, "paid": paid,
        "pct": paid / overdue if overdue else 0.0,
        "statements": sum(1 for r in rows if r["statement"]), "sent": sum(1 for r in rows if r["sent"]),
        "missing": sum(1 for r in rows if r["status"] == "كشف ناقص"),
        "matched": sum(1 for r in rows if r["matched"] == 1), "not_matched": sum(1 for r in rows if r["matched"] == 0),
        "not_returned": sum(1 for r in rows if r["matched"] is None),
        "exited": sum(1 for r in rows if r["exited"]), "stop_sale": sum(1 for r in rows if r["stop_sale"]),
        "negotiating": sum(1 for r in rows if r["negotiating"]), "excluded": sum(1 for r in rows if r["excluded"]),
        "carried": sum(1 for r in rows if r["carried"]),
        "cancelled": sum(1 for r in rows for l in r["letters"] if l["status"] == "ملغي"),
    }
    for name, kinds in LETTER_GROUPS:
        s[name] = sum(1 for r in rows for l in r["active_letters"] if l["kind"] in kinds)
    return s


def by_rep(rows):
    groups = {}
    for r in rows:
        groups.setdefault(r["rep"], []).append(r)
    out = [{"rep": rep, **summarize(rs)} for rep, rs in groups.items()]
    out.sort(key=lambda x: -x["overdue"])
    return out


def export_excel(db, cycle, path, rep=None):
    """Month-end report: summary, reps, customers, letters, receipts."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    from .letters import fmt_date
    rows = customer_rows(db, cycle)
    if rep:
        rows = [r for r in rows if r["rep"] == rep]
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="12355B")
    money_fmt = "#,##0.00"
    wb = openpyxl.Workbook()

    def sheet(title, headers, widths, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, w in enumerate(widths):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
        ws.freeze_panes = "A2"
        return ws

    def fmt(ws, money_cols=(), pct_cols=()):
        for row in ws.iter_rows(min_row=2):
            for c in money_cols:
                row[c].number_format = money_fmt
            for c in pct_cols:
                row[c].number_format = "0%"

    # summary
    s = summarize(rows)
    ws = sheet("الملخص", ["البيان", "القيمة"], (42, 22), first=True)
    title = f"دورة {cycle['month']}" + (f" — {rep}" if rep else "")
    for label, key, kind in (
            (title, None, None), ("عدد العملاء المطلوب مطابقتهم", "customers", "n"), ("الرصيد يوم 1", "balance", "m"),
            ("المتأخر يوم 1", "overdue", "m"), ("المسدد خلال الدورة", "paid", "m"), ("نسبة التحصيل من المتأخر", "pct", "p"),
            ("كشوف جاهزة", "statements", "n"), ("كشوف أرسلت للمناديب", "sent", "n"), ("كشوف ناقصة", "missing", "n"),
            ("مطابق", "matched", "n"), ("غير مطابق", "not_matched", "n"), ("لم يرجع الكشف", "not_returned", "n"),
            ("خرجوا من الدورة (سددوا النسبة)", "exited", "n"), ("خطابات الإدارة المالية 1003", "1003", "n"),
            ("خطابات قانونية 1004", "1004", "n"), ("إجراء قانوني 1005/1006", "1005/1006", "n"),
            ("خطابات ملغاة", "cancelled", "n"), ("إيقاف بيع", "stop_sale", "n"), ("تحت التفاوض", "negotiating", "n"),
            ("مستبعدون", "excluded", "n"), ("مستمرون من دورة سابقة", "carried", "n")):
        ws.append([label, s[key] if key else ""])
        cell = ws.cell(ws.max_row, 2)
        cell.number_format = {"m": money_fmt, "p": "0%"}.get(kind, "0")
    ws["A2"].font = Font(bold=True, size=13)

    # reps
    ws = sheet("المناديب", ["المندوب", "العملاء", "المتأخر يوم 1", "المسدد", "النسبة", "مطابق", "غير مطابق",
                            "لم يرجع", "خرجوا", "1003", "1004", "1005/1006", "إيقاف بيع"],
               (30, 9, 15, 15, 9, 9, 10, 9, 9, 8, 8, 10, 10))
    for r in by_rep(rows):
        ws.append([r["rep"], r["customers"], r["overdue"], r["paid"], r["pct"], r["matched"], r["not_matched"],
                   r["not_returned"], r["exited"], r["1003"], r["1004"], r["1005/1006"], r["stop_sale"]])
    fmt(ws, (2, 3), (4,))

    # customers
    ws = sheet("العملاء", ["رقم العميل", "اسم العميل", "المندوب", "النوع", "الرصيد يوم 1", "المتأخر يوم 1",
                           "رصيد الكشف", "أرسل الكشف", "المطابقة", "المسدد", "النسبة", "الحالة", "آخر خطاب",
                           "تاريخه", "كل الخطابات", "مستبعد/ملاحظات", "التنبيهات"],
               (12, 38, 24, 12, 14, 14, 14, 10, 11, 14, 8, 22, 22, 11, 40, 30, 50))
    for r in rows:
        last = r["last_letter"]
        ws.append([r["code"], r["name"], r["rep"], r["types"], r["balance"], r["overdue"], r["stmt_balance"],
                   "نعم" if r["sent"] else "", r["matched_text"], r["paid_month"], r["pct_month"], r["status"],
                   last["serial"] if last else "", fmt_date(last["date"]) if last else "",
                   " | ".join(f"{l['serial']}{' (ملغي)' if l['status'] == 'ملغي' else ''}" for l in r["letters"]),
                   " | ".join(x for x in (r["exclude_reason"], r["carried"],
                                          "تحت التفاوض" if r["negotiating"] else "",
                                          "إيقاف بيع" if r["stop_sale"] else "") if x),
                   r["warning"]])
    fmt(ws, (4, 5, 6, 9), (10,))

    # letters
    codes = {r["code"] for r in rows}
    ws = sheet("الخطابات", ["رقم الخطاب", "النوع", "التاريخ", "رقم العميل", "اسم العميل", "المندوب", "المبلغ",
                            "الحالة", "سبب الإلغاء", "الخطاب السابق", "المستخدم"],
               (24, 26, 11, 12, 36, 22, 14, 9, 30, 24, 14))
    for l in db.letters(cycle_id=cycle["id"]):
        if l["code"] in codes:
            ws.append([l["serial"], stages.KIND_NAMES.get(l["kind"], l["kind"]), fmt_date(l["date"]), l["code"],
                       l["name"], l["rep_name"], l["amount"], l["status"], l["cancel_reason"], l["prev_serial"],
                       l["user"]])
    fmt(ws, (6,))

    # receipts
    ws = sheet("سندات القبض", ["رقم السند", "التاريخ", "رقم العميل", "الاسم", "الحساب", "المبلغ"],
               (12, 11, 12, 36, 12, 14))
    for r in db.receipts(cycle["id"]):
        if r["code"] in codes:
            ws.append([r["voucher"], fmt_date(r["date"]), r["code"], r["name"], r["account"], r["amount"]])
    fmt(ws, (5,))
    wb.save(path)
    return path
