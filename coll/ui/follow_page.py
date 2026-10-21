"""Follow-up screen — Day 7 / 14 / 21: signed statements, receipts, the payment rule and the letters."""
import os

from PySide6.QtCore import QDate, QSettings, Qt, QUrl
from PySide6.QtGui import QColor, QDesktopServices
from PySide6.QtWidgets import (QAbstractItemView, QApplication, QComboBox, QDateEdit, QDialog, QDialogButtonBox,
                               QDoubleSpinBox, QFileDialog, QFormLayout, QGridLayout, QHBoxLayout, QInputDialog,
                               QLabel, QLineEdit, QMenu, QPushButton, QTableWidgetItem, QVBoxLayout,
                               QWidget)

from .. import mailer, receipts, stages
from .cycle_page import step_box
from .mail_preview import MailPreviewDialog
from .widgets import (NAVY, ColumnMapDialog, Progress, ask, color_item, error, info, money, month_label, ro, styled_table,
                      warn)

STATUS_COLORS = {"خطاب مالي": "#b54708", "خطاب قانوني — إيقاف بيع": "#b42318", "إحالة السند للتنفيذ": "#7a271a",
                 "إنذار دعوى": "#7a271a", "تحت التفاوض": "#6941c6", "مستبعد": "#667085", "مطابق": "#067647",
                 "غير مطابق": "#b42318"}
# letter types that can replace each other for one customer
SWITCH = {"1004": "1004PN", "1004PN": "1004", "1005": "1006", "1006": "1005"}
COLS = ["✔", "رقم العميل", "اسم العميل", "المندوب", "المتأخر يوم 1", "المسدد", "النسبة", "المطابقة",
        "الحالة", "الخطاب المقترح", "مبلغ الخطاب", "السبب / ملاحظات"]


def pct_text(p):
    return f"{(p or 0) * 100:.0f}%"


class NegotiationDialog(QDialog):
    """Form 1002: settlement / scheduling agreement."""

    def __init__(self, parent, name, amount):
        super().__init__(parent)
        self.setWindowTitle(f"اتفاقية تسوية (1002) — {name}")
        f = QFormLayout(self)
        self.amount = QDoubleSpinBox()
        self.amount.setRange(0, 1e10)
        self.amount.setDecimals(2)
        self.amount.setGroupSeparatorShown(True)
        self.amount.setValue(amount)
        self.rep, self.title, self.id_no = QLineEdit(), QLineEdit(), QLineEdit()
        self.option = QComboBox()
        self.option.addItems(["الخيار الأول: أقساط", "الخيار الثاني: سحب نقدي مع 130%"])
        f.addRow("مبلغ المديونية:", self.amount)
        f.addRow("ممثل العميل:", self.rep)
        f.addRow("صفته:", self.title)
        f.addRow("رقم الهوية:", self.id_no)
        f.addRow("آلية التسوية:", self.option)
        grid = QGridLayout()
        self.inst = []
        today = QDate.currentDate()
        for i in range(4):
            a = QDoubleSpinBox()
            a.setRange(0, 1e10)
            a.setDecimals(2)
            a.setGroupSeparatorShown(True)
            d = QDateEdit(today.addDays(7 * (i + 1)))
            d.setCalendarPopup(True)
            d.setDisplayFormat("dd/MM/yyyy")
            grid.addWidget(QLabel(f"القسط {i + 1}:"), i, 0)
            grid.addWidget(a, i, 1)
            grid.addWidget(QLabel("بتاريخ"), i, 2)
            grid.addWidget(d, i, 3)
            self.inst.append((a, d))
        b = QPushButton("توزيع المبلغ على 4 أقساط متساوية")
        b.clicked.connect(self._split)
        f.addRow(grid)
        f.addRow(b)
        note = QLabel("أثناء التفاوض يُستثنى العميل من المراحل، وينبهك البرنامج إذا فات موعد قسط دون سداد.")
        note.setWordWrap(True)
        f.addRow(note)
        bb = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        bb.button(QDialogButtonBox.Ok).setText("إصدار الاتفاقية")
        bb.accepted.connect(self.accept)
        bb.rejected.connect(self.reject)
        f.addRow(bb)

    def _split(self):
        part = round(self.amount.value() / 4, 2)
        for i, (a, _) in enumerate(self.inst):
            a.setValue(part if i < 3 else round(self.amount.value() - 3 * part, 2))

    def data(self):
        inst = [(a.value(), d.date().toString("yyyy-MM-dd")) for a, d in self.inst if a.value() > 0] \
            if self.option.currentIndex() == 0 else []
        return {"amount": self.amount.value(), "rep": self.rep.text().strip(), "title": self.title.text().strip(),
                "id_no": self.id_no.text().strip(), "option": self.option.currentIndex() + 1, "inst": inst}


class FollowPage(QWidget):
    def __init__(self, win):
        super().__init__()
        self.win = win
        self.cycle = None
        self.rows = []
        self.packages = []
        self.kind_override = {}   # (cycle, code) → letter type chosen by the user for this issue
        lay = QVBoxLayout(self)
        top = QHBoxLayout()
        top.addWidget(QLabel("الدورة:"))
        self.cycle_box = QComboBox()
        self.cycle_box.setMinimumWidth(180)
        self.cycle_box.currentIndexChanged.connect(self._cycle_changed)
        top.addWidget(self.cycle_box)
        self.alert = QLabel()
        self.alert.setStyleSheet("color:#b42318; font-weight:700;")
        top.addWidget(self.alert, 1)
        lay.addLayout(top)

        steps = QHBoxLayout()
        b1, v1 = step_box("١) الكشوف الموقعة (المطابقات)")
        self.signed_info = QLabel("سمِّ كل ملف برقم العميل (مثل 18010274.pdf) وارفعها مرة واحدة. أو اضغط كليك يمين "
                                  "على العميل في الجدول ← مطابق / غير مطابق.")
        self.signed_info.setWordWrap(True)
        v1.addWidget(self.signed_info)
        b = QPushButton("📥 رفع الكشوف الموقعة")
        b.clicked.connect(self.upload_signed)
        v1.addWidget(b)
        v1.addStretch()
        b2, v2 = step_box("٢) سندات القبض")
        self.rec_info = QLabel("ارفع تقرير سندات القبض من النظام (PDF أو Excel). يُحسب فقط ما له رقم عميل في "
                               "الحساب التحليلي، وتُستبعد التحويلات بين الصناديق والبنوك.")
        self.rec_info.setWordWrap(True)
        v2.addWidget(self.rec_info)
        for text, fn in (("📥 رفع سندات قبض", self.upload_receipts), ("🧾 عرض السندات المرفوعة", self.show_receipts)):
            b = QPushButton(text)
            b.clicked.connect(fn)
            v2.addWidget(b)
        per = QGridLayout()
        self.p_start = QDateEdit()
        self.p_end = QDateEdit()
        for d in (self.p_start, self.p_end):
            d.setCalendarPopup(True)
            d.setDisplayFormat("dd/MM/yyyy")
        self.p_start.setToolTip("أول يوم في فترة التحصيل لهذه الدورة (يُحفظ)")
        self.p_end.setToolTip("آخر يوم يُحسب للمرحلة المختارة — يتغير تلقائيًا مع المرحلة")
        per.addWidget(QLabel("الفترة من:"), 0, 0)
        per.addWidget(self.p_start, 0, 1)
        per.addWidget(QLabel("إلى:"), 0, 2)
        per.addWidget(self.p_end, 0, 3)
        v2.addLayout(per)
        self.p_start.dateChanged.connect(self._start_changed)
        self.p_end.dateChanged.connect(self._end_changed)
        v2.addStretch()
        b3, v3 = step_box("٣) الخطابات")
        row = QHBoxLayout()
        self.stage = QComboBox()
        for key, (lvl, name) in stages.STAGES.items():
            self.stage.addItem(name, key)
        self.stage.currentIndexChanged.connect(self._stage_changed)
        self.date = QDateEdit(QDate.currentDate())
        self.date.setCalendarPopup(True)
        self.date.setDisplayFormat("dd/MM/yyyy")
        row.addWidget(QLabel("المرحلة:"))
        row.addWidget(self.stage, 1)
        row.addWidget(QLabel("تاريخ الخطاب:"))
        row.addWidget(self.date)
        v3.addLayout(row)
        self.letter_info = QLabel()
        self.letter_info.setWordWrap(True)
        v3.addWidget(self.letter_info)
        grid = QGridLayout()
        for i, (text, fn) in enumerate((("🖨 إصدار خطابات المحددين", self.issue),
                                        ("📂 فتح فولدر الخطابات", self.open_letters),
                                        ("✉ إرسال الخطابات للمناديب", self.send),
                                        ("📊 تصدير الجدول Excel", self.export))):
            b = QPushButton(text)
            b.clicked.connect(fn)
            grid.addWidget(b, i // 2, i % 2)
        v3.addLayout(grid)
        v3.addStretch()
        for b in (b1, b2, b3):
            steps.addWidget(b, 1)
        lay.addLayout(steps)

        bar = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText("بحث بالاسم أو الرقم...")
        self.search.textChanged.connect(self.fill)
        self.view = QComboBox()
        self.view.addItems(["المرشحون للمرحلة المختارة", "كل العملاء المطلوب مطابقتهم"])
        self.view.currentIndexChanged.connect(self.refresh)
        self.f_rep = QComboBox()
        self.f_rep.currentIndexChanged.connect(self.fill)
        bar.addWidget(self.search, 2)
        bar.addWidget(QLabel("عرض:"))
        bar.addWidget(self.view, 1)
        bar.addWidget(QLabel("المندوب:"))
        bar.addWidget(self.f_rep, 1)
        for text, val in (("تحديد الكل", True), ("إلغاء التحديد", False)):
            b = QPushButton(text)
            b.clicked.connect(lambda _=None, v=val: self.check_all(v))
            bar.addWidget(b)
        lay.addLayout(bar)
        self.summary = QLabel()
        self.summary.setStyleSheet(f"color:{NAVY}; font-weight:700;")
        lay.addWidget(self.summary)
        self.table = styled_table(COLS)
        self.table.setEditTriggers(QAbstractItemView.DoubleClicked | QAbstractItemView.EditKeyPressed)
        self.table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self.menu)
        self.table.cellDoubleClicked.connect(self._double)
        lay.addWidget(self.table, 1)

    # ---------------------------------------------------------------- data
    def load(self):
        db = self.win.db
        cur = self.cycle["month"] if self.cycle else None
        self.cycle_box.blockSignals(True)
        self.cycle_box.clear()
        for c in db.cycles():
            self.cycle_box.addItem(f"دورة {month_label(c['month'])}", c["month"])
        if cur:
            self.cycle_box.setCurrentIndex(max(0, self.cycle_box.findData(cur)))
        self.cycle_box.blockSignals(False)
        self._cycle_changed()

    def _cycle_changed(self):
        m = self.cycle_box.currentData()
        self.cycle = self.win.db.cycle(m) if m else None
        if self.cycle:
            # suggest the stage that is due today
            lvl = max((c["level"] or 0) for c in self.win.db.cycle_customers(self.cycle["id"], True)) \
                if self.win.db.cycle_customers(self.cycle["id"], True) else 0
            self.stage.blockSignals(True)
            self.stage.setCurrentIndex(min(lvl, 2))
            self.stage.blockSignals(False)
            self._set_period()
        self.refresh()

    def stage_key(self):
        return self.stage.currentData()

    # ---------------------------------------------------------------- receipts period
    def period(self):
        return self.p_start.date().toString("yyyy-MM-dd"), self.p_end.date().toString("yyyy-MM-dd")

    def _set_period(self):
        """Start from the cycle, end from the selected stage; recompute paid for that period."""
        if not self.cycle:
            return
        db = self.win.db
        for w in (self.p_start, self.p_end):
            w.blockSignals(True)
        self.p_start.setDate(QDate.fromString(stages.cycle_start(db, self.cycle), "yyyy-MM-dd"))
        self.p_end.setDate(QDate.fromString(stages.stage_end(db, self.cycle, self.stage_key()), "yyyy-MM-dd"))
        for w in (self.p_start, self.p_end):
            w.blockSignals(False)
        stages.recompute_paid(db, self.cycle, *self.period())

    def _stage_changed(self):
        self._set_period()
        self.refresh()

    def _start_changed(self):
        if self.cycle and self.win.can_edit(silent=True):
            self.win.db.set_cycle_start(self.cycle["id"], self.p_start.date().toString("yyyy-MM-dd"))
            self._set_period()
            self.refresh()

    def _end_changed(self):
        if self.cycle:
            stages.recompute_paid(self.win.db, self.cycle, *self.period())
            self.refresh()

    def refresh(self):
        db = self.win.db
        self.packages = []
        if not self.cycle:
            self.rows = []
            self.alert.setText("")
            self.letter_info.setText("ابدأ الدورة من شاشة «الدورة الشهرية».")
            self.fill()
            return
        key = self.stage_key()
        if self.view.currentIndex() == 0:
            self.rows = stages.candidates(db, self.cycle, key)
            for r in self.rows:
                k = self.kind_override.get((self.cycle["id"], r["code"]))
                if k and SWITCH.get(r["kind"]) == k:
                    r["kind"] = k
                    r["reason"] = (r["reason"] + "، نوع الخطاب مختار يدويًا").strip("، ")
        else:
            thr = stages.threshold(db, key)
            self.rows = []
            masters = db.customers()
            for c in db.cycle_customers(self.cycle["id"], required_only=True):
                r = dict(c)
                r.update(eligible=False, kind="", amount=stages.remaining(c, db.setting("letter_amount")),
                         reason=c["exclude_reason"] or c["carried"] or "",
                         has_pn=masters.get(c["code"], {}).get("has_pn"))
                if (c["pct"] or 0) >= thr:
                    r["reason"] = (r["reason"] + " خرج من الدورة (سدد النسبة)").strip()
                self.rows.append(r)
        reps = sorted({r["rep_name"] for r in self.rows if r["rep_name"]})
        cur = self.f_rep.currentText()
        self.f_rep.blockSignals(True)
        self.f_rep.clear()
        self.f_rep.addItem("الكل")
        self.f_rep.addItems(reps)
        self.f_rep.setCurrentIndex(max(0, self.f_rep.findText(cur)))
        self.f_rep.blockSignals(False)
        stop = stages.stop_sale(db, self.cycle)
        late = stages.late_instalments(db)
        self.alert.setText("   ".join(x for x in (
            f"⛔ إيقاف بيع: {len(stop)} عميل" if stop else "",
            f"⏰ أقساط تفاوض متأخرة: {len(late)}" if late else "") if x))
        eligible = [r for r in self.rows if r["eligible"]]
        lvl, name = stages.STAGES[key]
        self.letter_info.setText(
            f"{name}: {len(eligible)} عميل مرشح بمبلغ {money(sum(r['amount'] for r in eligible))}. "
            f"نسبة الإعفاء في هذه المرحلة: {stages.threshold(db, key) * 100:.0f}% من المتأخر يوم 1." if self.view.currentIndex() == 0
            else "عرض كل العملاء — اختر «المرشحون» لإصدار الخطابات.")
        self.fill()

    def visible_rows(self):
        s = self.search.text().strip()
        rep = self.f_rep.currentText()
        out = []
        for r in self.rows:
            if s and s not in r["code"] and s not in (r["name"] or ""):
                continue
            if rep and rep != "الكل" and r["rep_name"] != rep:
                continue
            out.append(r)
        return out

    def fill(self):
        rows = self.visible_rows()
        t = self.table
        t.blockSignals(True)
        t.setRowCount(len(rows))
        candidates_view = self.view.currentIndex() == 0
        for i, r in enumerate(rows):
            chk = QTableWidgetItem()
            if candidates_view:
                chk.setFlags(Qt.ItemIsUserCheckable | Qt.ItemIsEnabled)
                chk.setCheckState(Qt.Checked if r["eligible"] else Qt.Unchecked)
            else:
                chk.setFlags(Qt.ItemIsEnabled)
            t.setItem(i, 0, chk)
            t.setItem(i, 1, ro(r["code"]))
            t.setItem(i, 2, ro(r["name"]))
            t.setItem(i, 3, ro(r["rep_name"]))
            t.setItem(i, 4, ro(money(r["overdue"]), True))
            t.setItem(i, 5, ro(money(r["paid"] or 0), True))
            it = ro(pct_text(r["pct"]), True)
            color_item(it, "#067647" if (r["pct"] or 0) >= stages.threshold(self.win.db, self.stage_key()) else "#b42318")
            t.setItem(i, 6, it)
            t.setItem(i, 7, ro({1: "✔ مطابق", 0: "✖ غير مطابق"}.get(r["matched"], "—")))
            it = ro(r["status"])
            color_item(it, STATUS_COLORS.get(r["status"], "#344054"))
            t.setItem(i, 8, it)
            t.setItem(i, 9, ro(stages.KIND_NAMES.get(r["kind"], "") + (f" ({r['kind'][:4]})" if r["kind"] else "")))
            amt = QTableWidgetItem(f"{r['amount']:.2f}")
            amt.setTextAlignment(Qt.AlignAbsolute | Qt.AlignRight | Qt.AlignVCenter)
            if not candidates_view:
                amt.setFlags(amt.flags() & ~Qt.ItemIsEditable)
            amt.setToolTip("اضغط مرتين لتعديل مبلغ الخطاب")
            t.setItem(i, 10, amt)
            t.setItem(i, 11, ro(r["reason"]))
            if not r["eligible"] and candidates_view:
                for c in range(t.columnCount()):
                    if t.item(i, c):
                        t.item(i, c).setBackground(QColor("#f2f4f7"))
        t.blockSignals(False)
        t.resizeColumnsToContents()
        t.setColumnWidth(2, min(t.columnWidth(2), 280))
        total_paid = sum(r["paid"] or 0 for r in rows)
        total_over = sum(r["overdue"] or 0 for r in rows)
        self.summary.setText(f"المعروض: {len(rows)} عميل   |   المتأخر يوم 1: {money(total_over)}   |   المسدد: "
                             f"{money(total_paid)} ({pct_text(total_paid / total_over if total_over else 0)})")

    def _row_at(self, i):
        code = self.table.item(i, 1).text()
        return next(r for r in self.rows if r["code"] == code)

    def _double(self, row, col):
        if col != 10 or self.view.currentIndex() != 0:
            self.win.open_customer(self.table.item(row, 1).text())
            self.refresh()

    def check_all(self, value):
        for i in range(self.table.rowCount()):
            it = self.table.item(i, 0)
            if it.flags() & Qt.ItemIsUserCheckable:
                it.setCheckState(Qt.Checked if value else Qt.Unchecked)

    def selected_items(self):
        out = []
        for i in range(self.table.rowCount()):
            if self.table.item(i, 0).checkState() == Qt.Checked:
                r = self._row_at(i)
                try:
                    amount = float(self.table.item(i, 10).text().replace(",", ""))
                except ValueError:
                    amount = r["amount"]
                out.append({"code": r["code"], "kind": r["kind"], "amount": amount, "name": r["name"],
                            "eligible": r["eligible"]})
        return out

    # ---------------------------------------------------------------- context menu
    def menu(self, pos):
        idx = self.table.indexAt(pos)
        if not idx.isValid() or not self.cycle:
            return
        r = self._row_at(idx.row())
        m = QMenu(self)
        acts = [("✔ مطابق", lambda: self._match(r, 1)), ("✖ غير مطابق", lambda: self._match(r, 0)),
                ("📎 إرفاق كشف موقع لهذا العميل", lambda: self._attach_one(r)), None]
        if r["excluded"]:
            acts.append(("↩ إلغاء الاستبعاد", lambda: self._include(r)))
        else:
            acts.append(("🚫 استبعاد من الإجراءات هذه الدورة (مع سبب)", lambda: self._exclude(r)))
            acts.append(("⛔ استثناء دائم (لا يدخل أي دورة)", lambda: self._exclude(r, permanent=True)))
        if self.view.currentIndex() == 0 and r["kind"] in SWITCH:
            other = SWITCH[r["kind"]]
            acts.append((f"🔁 تغيير نوع الخطاب إلى: {stages.KIND_NAMES[other]} ({other[:4]})",
                         lambda: self._switch(r, other)))
        acts += [("🤝 بدء تفاوض (اتفاقية 1002)", lambda: self.negotiate(r)), None,
                 ("👤 فتح بطاقة العميل", lambda: (self.win.open_customer(r["code"]), self.refresh()))]
        for a in acts:
            if a is None:
                m.addSeparator()
            else:
                m.addAction(a[0]).triggered.connect(a[1])
        m.exec(self.table.viewport().mapToGlobal(pos))

    def _match(self, r, value):
        if self.win.can_edit():
            stages.set_matched(self.win.db, self.cycle, r["code"], value)
            self.refresh()

    def _attach_one(self, r):
        if not self.win.can_edit():
            return
        path, _ = QFileDialog.getOpenFileName(self, f"الكشف الموقع — {r['name']}",
                                              QSettings().value("last_dir", os.path.expanduser("~")),
                                              "PDF / صور (*.pdf *.jpg *.jpeg *.png)")
        if path:
            stages.attach_signed_to(self.win.db, self.win.folder, self.cycle, r["code"], path)
            self.refresh()

    def _exclude(self, r, permanent=False):
        if not self.win.can_edit():
            return
        title = "استثناء دائم" if permanent else "استبعاد من الإجراءات"
        reason, ok = QInputDialog.getText(self, title, f"سبب {title} — {r['name']}:")
        if ok and reason.strip():
            if permanent:
                db = self.win.db
                c = db.customer(r["code"])
                c.update(perm_exclude=1, perm_reason=reason.strip())
                db.save_customer(c)
                db.log("استثناء دائم", reason.strip(), code=r["code"])
            stages.exclude(self.win.db, self.cycle, r["code"], ("استثناء دائم: " if permanent else "") + reason.strip())
            self.refresh()
        elif ok:
            warn(self, "لازم تكتب السبب.")

    def _switch(self, r, kind):
        self.kind_override[(self.cycle["id"], r["code"])] = kind
        self.refresh()

    def _include(self, r):
        if self.win.can_edit():
            stages.include(self.win.db, self.cycle, r["code"])
            self.refresh()

    def negotiate(self, r):
        if not self.win.can_edit():
            return
        dlg = NegotiationDialog(self, r["name"], stages.remaining(r, self.win.db.setting("letter_amount")))
        if dlg.exec() != QDialog.Accepted:
            return
        try:
            pdf = stages.start_negotiation(self.win.db, self.win.folder, self.cycle, r["code"], dlg.data())
        except Exception as e:  # noqa: BLE001
            error(self, f"تعذر إصدار الاتفاقية:\n{e}")
            return
        self.refresh()
        QDesktopServices.openUrl(QUrl.fromLocalFile(pdf))

    # ---------------------------------------------------------------- step 1
    def upload_signed(self):
        if not self.win.can_edit() or not self.cycle:
            return
        paths, _ = QFileDialog.getOpenFileNames(self, "الكشوف الموقعة", QSettings().value("last_dir", os.path.expanduser("~")),
                                                "PDF / صور (*.pdf *.jpg *.jpeg *.png)")
        if not paths:
            return
        QSettings().setValue("last_dir", os.path.dirname(paths[0]))
        matched, unknown = stages.attach_signed(self.win.db, self.win.folder, self.cycle, paths)
        self.refresh()
        msg = f"تم تسجيل {len(matched)} عميل «مطابق»."
        if unknown:
            msg += (f"\n\n{len(unknown)} ملف بدون رقم عميل معروف في اسمه (أرفقها من كليك يمين على العميل):\n"
                    + "\n".join(unknown[:15]))
        info(self, msg)

    # ---------------------------------------------------------------- step 2
    def upload_receipts(self):
        if not self.win.can_edit() or not self.cycle:
            return
        paths, _ = QFileDialog.getOpenFileNames(self, "سندات القبض", QSettings().value("last_dir", os.path.expanduser("~")),
                                                "PDF / Excel (*.pdf *.xlsx *.xlsm *.csv)")
        if not paths:
            return
        QSettings().setValue("last_dir", os.path.dirname(paths[0]))
        db = self.win.db
        known = set(db.customers())
        lines = []
        for p in paths:
            try:
                if p.lower().endswith(".pdf"):
                    QApplication.setOverrideCursor(Qt.WaitCursor)
                    try:
                        recs, excluded = receipts.read_pdf(p, known)
                    finally:
                        QApplication.restoreOverrideCursor()
                else:
                    recs, excluded = self._read_excel(p, known)
                    if recs is None:
                        continue
            except Exception as e:  # noqa: BLE001
                lines.append(f"⛔ {os.path.basename(p)}: تعذر القراءة ({e})")
                continue
            start, end = self.period()
            res = stages.add_receipts(db, self.win.folder, self.cycle, recs, p, start, end)
            if res["before"] or res["after"]:
                lines.append(f"📅 الفترة المقبولة {self.p_start.date().toString('dd/MM/yyyy')} — "
                             f"{self.p_end.date().toString('dd/MM/yyyy')}: "
                             + (f"رُفض {res['before']} سند قبل بداية الدورة. " if res["before"] else "")
                             + (f"رُفض {res['after']} سند بعد نهاية المرحلة — ارفعها مع المرحلة القادمة."
                                if res["after"] else ""))
            lines.append(f"🧾 {os.path.basename(p)}: {res['count']} سند عميل في الفترة ({res['added']} جديد"
                         + (f"، {res['duplicates']} مكرر لم يُحسب مرتين" if res["duplicates"] else "") + ")، "
                         f"استُبعد {excluded} سطر تحويلات داخلية. المحسوب للعملاء المطلوبين: {money(res['total'])}")
            if res["outside"]:
                lines.append(f"   ℹ {len(res['outside'])} عميل سدد وليس عليه مطابقة هذه الدورة.")
        self.refresh()
        info(self, "\n".join(lines))

    def _read_excel(self, path, known):
        from ..aging import read_table
        db = self.win.db
        rows = read_table(path)
        saved = db.column_map("receipts")
        first, mapping = receipts.guess_columns(rows)
        if saved:
            first, mapping = saved["first"], {k: int(v) for k, v in saved["mapping"].items()}
        if not saved or "code" not in mapping or "amount" not in mapping:
            dlg = ColumnMapDialog(self, rows, receipts.ROLES, receipts.ROLE_LABELS, first, mapping,
                                  "حدد عمود رقم العميل (الحساب التحليلي) والمبلغ على الأقل. يُحفظ الإعداد للمرات القادمة.")
            if dlg.exec() != QDialog.Accepted:
                return None, 0
            mapping, first = dlg.mapping(), dlg.first_row()
            if "code" not in mapping or "amount" not in mapping:
                warn(self, "لازم تحدد عمود رقم العميل وعمود المبلغ.")
                return None, 0
            db.save_column_map("receipts", {"mapping": mapping, "first": first})
        return receipts.read_rows(rows, first, mapping, known)

    def show_receipts(self):
        if not self.cycle:
            return
        db = self.win.db
        recs = db.receipts(self.cycle["id"])
        dlg = QDialog(self)
        dlg.setWindowTitle(f"سندات القبض — دورة {month_label(self.cycle['month'])}")
        dlg.resize(1000, 600)
        v = QVBoxLayout(dlg)
        t = styled_table(["رقم السند", "التاريخ", "رقم العميل", "الاسم", "الحساب", "المبلغ", "الملف"])
        t.setRowCount(len(recs))
        for i, r in enumerate(recs):
            for c, val in enumerate((r["voucher"], r["date"], r["code"], r["name"], r["account"])):
                t.setItem(i, c, ro(val))
            t.setItem(i, 5, ro(money(r["amount"]), True))
            t.setItem(i, 6, ro(os.path.basename(r["source"] or "")))
        t.resizeColumnsToContents()
        v.addWidget(QLabel(f"{len(recs)} سند — الإجمالي {money(sum(r['amount'] for r in recs))}"))
        v.addWidget(t, 1)
        files = sorted({r["source"] for r in recs if r["source"]})
        if files:
            bar = QHBoxLayout()
            cb = QComboBox()
            for f in files:
                cb.addItem(os.path.basename(f), f)
            b = QPushButton("🗑 حذف سندات هذا الملف")

            def remove():
                if self.win.can_edit() and ask(dlg, f"حذف كل سندات الملف {cb.currentText()} من الدورة؟"):
                    db.delete_receipts(self.cycle["id"], cb.currentData())
                    db.log("حذف سندات قبض", cb.currentText(), cycle_id=self.cycle["id"])
                    stages.recompute_paid(db, self.cycle, *self.period())
                    dlg.accept()
                    self.refresh()
            b.clicked.connect(remove)
            bar.addWidget(QLabel("ملف:"))
            bar.addWidget(cb, 1)
            bar.addWidget(b)
            v.addLayout(bar)
        dlg.exec()

    # ---------------------------------------------------------------- step 3
    def issue(self):
        if not self.win.can_edit() or not self.cycle:
            return
        if self.view.currentIndex() != 0:
            warn(self, "اختر عرض «المرشحون للمرحلة المختارة» أولًا.")
            return
        items = self.selected_items()
        if not items:
            warn(self, "حدد العملاء المطلوب إصدار خطابات لهم.")
            return
        forced = [i for i in items if not i["eligible"]]
        msg = f"إصدار {len(items)} خطاب ({self.stage.currentText()}) بتاريخ {self.date.date().toString('dd/MM/yyyy')}؟"
        if forced:
            msg += f"\n\n⚠ منهم {len(forced)} عميل غير مرشح (سدد/مستبعد/تفاوض) وحددته يدويًا."
        if not ask(self, msg):
            return
        db = self.win.db
        if not db.setting("contact_mobile") and not ask(self, "جوال/إيميل التواصل في الخطابات فاضي (من شاشة النماذج). "
                                                              "تكمل بدونها؟"):
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        bar = Progress(self, "إصدار الخطابات")
        try:
            self.packages = stages.issue(db, self.win.folder, self.cycle, items,
                                         self.date.date().toString("yyyy-MM-dd"), self.stage_key(), bar)
        except Exception as e:  # noqa: BLE001
            bar.close()
            QApplication.restoreOverrideCursor()
            error(self, f"تعذر إصدار الخطابات:\n{e}")
            self.refresh()
            return
        bar.close()
        QApplication.restoreOverrideCursor()
        packages = self.packages
        self.refresh()
        self.packages = packages
        info(self, "تم إصدار الخطابات:\n" + "\n".join(
            f"• {p['rep']}: {len(p['letters'])} خطاب" + (f" ← {p['email']}" if p["email"] else " (بدون إيميل)")
            for p in packages) + "\n\nلكل مندوب ملف PDF واحد للطباعة وملف مضغوط في فولدر الخطابات.")

    def open_letters(self):
        if self.cycle:
            QDesktopServices.openUrl(QUrl.fromLocalFile(self.win.folder.files_dir(self.cycle["month"], "letters")))

    def send(self):
        if not self.win.can_edit():
            return
        if not self.packages:
            warn(self, "أصدر الخطابات أولًا — يُرسل آخر إصدار فقط.")
            return
        cfg = self.win.mail_config()
        if not cfg:
            warn(self, "اضبط إعدادات الإيميل أولًا من شاشة الإعدادات، أو أرسل الملفات يدويًا من فولدر الخطابات.")
            return
        month = month_label(self.cycle["month"])
        stage = self.stage.currentText()
        mails = [{"rep": p["rep"], "to": p["email"], "subject": f"{stage} — عملاء {p['rep']} — {month}",
                  "body": (f"الأستاذ/ {p['rep']} المحترم،\n\nمرفق {stage} لعملائك ({len(p['letters'])} عميل) لشهر "
                           f"{month}.\nنرجو تسليمها للعملاء ومتابعة السداد.\n\nمع الشكر،\n"
                           "الإدارة المالية — قسم التحصيل والائتمان"),
                  "files": [p["zip"]], "count": f"{len(p['letters'])} خطاب"} for p in self.packages]
        dlg = MailPreviewDialog(self, mails, cfg.get("cc", ""), "مراجعة إيميلات الخطابات قبل الإرسال")
        if dlg.exec() != QDialog.Accepted:
            return
        msgs = [(m["to"], m["subject"], m["body"], m["files"]) for m in dlg.selected()]
        if not msgs:
            return
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            results = mailer.send_all(cfg, msgs)
        except Exception as e:  # noqa: BLE001
            QApplication.restoreOverrideCursor()
            error(self, f"تعذر الاتصال بخادم الإيميل:\n{e}")
            return
        QApplication.restoreOverrideCursor()
        ok = [r for r in results if r[1]]
        self.win.db.log("إرسال خطابات", f"{stage}: {len(ok)} من {len(results)}", cycle_id=self.cycle["id"])
        info(self, "\n".join(f"{'✔' if r[1] else '✖'} {r[0]}" + ("" if r[1] else f" — {r[2]}") for r in results))

    def export(self):
        if not self.cycle:
            return
        path, _ = QFileDialog.getSaveFileName(self, "تصدير", os.path.join(os.path.expanduser("~"),
                                              f"متابعة التحصيل {self.cycle['month']}.xlsx"), "Excel (*.xlsx)")
        if not path:
            return
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.sheet_view.rightToLeft = True
        ws.append(COLS[1:])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="12355B")
        for r in self.visible_rows():
            ws.append([r["code"], r["name"], r["rep_name"], r["overdue"], r["paid"] or 0, r["pct"] or 0,
                       {1: "مطابق", 0: "غير مطابق"}.get(r["matched"], ""), r["status"],
                       stages.KIND_NAMES.get(r["kind"], ""), r["amount"], r["reason"]])
        for row in ws.iter_rows(min_row=2):
            for c in (3, 4, 9):
                row[c].number_format = "#,##0.00"
            row[5].number_format = "0%"
        for col, w in zip("ABCDEFGHIJK", (12, 40, 20, 14, 14, 9, 12, 22, 22, 14, 40)):
            ws.column_dimensions[col].width = w
        wb.save(path)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))
