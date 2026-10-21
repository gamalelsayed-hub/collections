"""Dashboard figures and the month-end Excel report (synthetic data)."""
import openpyxl

from coll import receipts, report, stages
from tests.test_stages import setup_cycle


def test_summary_and_excel(tmp_path, qapp):
    folder, db, cyc = setup_cycle(tmp_path)
    f = tmp_path / "r.pdf"
    f.write_bytes(b"x")
    stages.add_receipts(db, folder, cyc, [receipts.Receipt("1", "2026-10-03", "1140101", "", "90000001", 400.0)],
                        str(f))
    db.update_cycle_customer(cyc["id"], "90000001", matched=1)
    stages.issue(db, folder, cyc, [{"code": "90000003", "kind": "1003", "amount": 50.0}], "2026-10-07")
    rows = report.customer_rows(db, cyc)
    s = report.summarize(rows)
    assert s["customers"] == 3 and s["overdue"] == 920 and s["paid"] == 400
    assert s["matched"] == 1 and s["exited"] == 1 and s["1003"] == 1
    reps = {r["rep"]: r for r in report.by_rep(rows)}
    assert reps["سالم عبدالله"]["customers"] == 2 and reps["خالد"]["1003"] == 1

    path = report.export_excel(db, cyc, str(tmp_path / "r.xlsx"))
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["الملخص", "المناديب", "العملاء", "الخطابات", "سندات القبض"]
    assert wb["الخطابات"].max_row == 2 and wb["سندات القبض"].max_row == 2 and wb["العملاء"].max_row == 4
    one = openpyxl.load_workbook(report.export_excel(db, cyc, str(tmp_path / "k.xlsx"), rep="خالد"))
    assert one["العملاء"].max_row == 2
