"""Customer attachments, matched customers without letters, data update form, direct confirmation and the rep's
match-status sheet (synthetic data)."""
import os

from coll import attachments as A
from coll import day1, direct, stages
from tests.test_stages import setup_cycle


def test_attachments(tmp_path):
    folder, db, cyc = setup_cycle(tmp_path)
    src = tmp_path / "سجل.pdf"
    src.write_bytes(b"%PDF-1.4 test")
    fid = A.add(db, folder, "90000001", "cr", str(src))
    recs = db.customer_files("90000001")
    assert len(recs) == 1 and recs[0]["id"] == fid and recs[0]["kind"] == "cr"
    p = A.full_path(folder, recs[0])
    assert os.path.exists(p) and p.startswith(os.path.join(folder.path, "files", "customers", "90000001"))
    cust = db.customer("90000001")
    assert "cr" not in A.missing(db, cust) and "vat" in A.missing(db, cust) and "pn" not in A.missing(db, cust)
    assert "pn" in A.missing(db, {**cust, "has_pn": 1})
    A.remove(db, folder, recs[0])
    assert not db.customer_files("90000001") and not os.path.exists(p)


def test_matched_customer_without_letters(tmp_path):
    folder, db, cyc = setup_cycle(tmp_path)
    stages.set_matched(db, cyc, "90000003", 1)
    assert {c["code"]: c for c in stages.candidates(db, cyc, "finance")}["90000003"]["eligible"]
    stages.set_stop_letters(db, cyc, "90000003", True, "عميل مهم")
    c = {c["code"]: c for c in stages.candidates(db, cyc, "finance")}["90000003"]
    assert not c["eligible"] and "عدم استكمال" in c["reason"]
    assert db.cycle_customer(cyc["id"], "90000003")["status"] == "مطابق — بدون خطابات"
    stages.set_stop_letters(db, cyc, "90000003", False)
    assert {c["code"]: c for c in stages.candidates(db, cyc, "finance")}["90000003"]["eligible"]
    # not matched any more → the decision is dropped
    stages.set_stop_letters(db, cyc, "90000003", True)
    stages.set_matched(db, cyc, "90000003", 0, "فاتورة غير صحيحة")
    c = db.cycle_customer(cyc["id"], "90000003")
    assert c["stop_letters"] == 0 and c["match_reason"] == "فاتورة غير صحيحة"


def test_update_form_and_direct_confirmation(tmp_path, qapp):
    folder, db, cyc = setup_cycle(tmp_path)
    codes = [c["code"] for c in direct.candidates(db, cyc)]
    assert "90000003" in codes
    pk = direct.issue_forms(db, folder, cyc, ["90000003"], "2026-10-20")
    assert pk and os.path.exists(pk[0]["zip"]) and pk[0]["letters"][0]["serial"] == "SL-26-1005-0001"
    c = db.cycle_customer(cyc["id"], "90000003")
    assert c["upd_serial"] == "SL-26-1005-0001" and db.letters(code="90000003")[0]["kind"] == "UPD"
    assert (db.cycle_customer(cyc["id"], "90000003")["level"] or 0) == 0       # not a stage letter
    # no e-mail / no statement yet → not sent
    mails, problems = direct.mails(db, folder, cyc, ["90000003"])
    assert not mails and "إيميل" in problems[0]
    cust = db.customer("90000003")
    cust["email"] = "client@example.com"
    db.save_customer(cust)
    stmt = os.path.join(folder.files_dir("2026-10", "statements"), "90000003.pdf")
    open(stmt, "wb").write(b"%PDF-1.4")
    db.update_cycle_customer(cyc["id"], "90000003", statement=os.path.relpath(stmt, folder.path))
    mails, problems = direct.mails(db, folder, cyc, ["90000003"], today="2026-10-20")
    assert not problems and mails[0]["to"] == "client@example.com" and mails[0]["files"] == [stmt]
    assert "7 أيام" in mails[0]["body"] and "27/10/2026" in mails[0]["body"]
    direct.mark_sent(db, cyc, "90000003", "client@example.com", today="2026-10-20")
    assert db.cycle_customer(cyc["id"], "90000003")["status"].startswith("مصادقة مباشرة")
    assert direct.apply_silent_acceptance(db, cyc, today="2026-10-26") == []
    assert direct.apply_silent_acceptance(db, cyc, today="2026-10-27") == [("90000003", "عميل ج")]
    c = db.cycle_customer(cyc["id"], "90000003")
    assert c["matched"] == 1 and "لعدم الرد" in c["match_reason"]
    # matched now: no longer in the list, and the letters stages continue (not paid)
    assert "90000003" not in [x["code"] for x in direct.candidates(db, cyc)]
    assert {x["code"]: x for x in stages.candidates(db, cyc, "finance")}["90000003"]["eligible"]


def test_rep_sheet_round_trip(tmp_path):
    folder, db, cyc = setup_cycle(tmp_path)
    custs = [dict(c, stmt_detail="", warning="") for c in db.cycle_customers(cyc["id"], required_only=True)]
    path = str(tmp_path / "قائمة العملاء.xlsx")
    day1._summary_sheet(custs, "سالم", "2026-10", path)
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws["I3"].value == "حالة المطابقة" and ws["J3"].value == "سبب عدم المطابقة"
    assert ws.data_validations.dataValidation and "مطابق" in ws.data_validations.dataValidation[0].formula1
    codes = [ws.cell(r, 1).value for r in range(4, 4 + len(custs))]
    ws.cell(4, 9, "مطابق")
    ws.cell(5, 9, "غير مطابق")
    ws.cell(5, 10, "فرق فاتورة 120")
    if len(custs) > 2:
        ws.cell(6, 9, "لم يرد")
    wb.save(path)
    entries = day1.read_rep_sheet(path)
    assert entries[0] == (str(codes[0]), "مطابق", "") and entries[1][1:] == ("غير مطابق", "فرق فاتورة 120")
    applied, unknown, invalid = day1.apply_rep_sheet(db, cyc, entries)
    assert applied == len(entries) and not unknown and not invalid
    assert db.cycle_customer(cyc["id"], str(codes[0]))["matched"] == 1
    c2 = db.cycle_customer(cyc["id"], str(codes[1]))
    assert c2["matched"] == 0 and c2["match_reason"] == "فرق فاتورة 120"
    if len(custs) > 2:
        c3 = db.cycle_customer(cyc["id"], str(codes[2]))
        assert c3["matched"] is None and c3["match_reason"] == "لم يرد"


def test_letter_paper(tmp_path, qapp):
    from PySide6.QtGui import QColor, QImage
    from coll import letters as L
    folder, db, cyc = setup_cycle(tmp_path)
    img = QImage(400, 566, QImage.Format_RGB32)
    img.fill(QColor("white"))
    src = str(tmp_path / "paper.png")
    img.save(src)
    page = L.page_image(src)
    assert page.width() == 1654 and page.height() == 2339
    page.save(folder.asset(stages.LETTERHEAD), "PNG")
    lh = stages.letterhead(db, folder)
    assert lh.page is not None and (lh.top, lh.bottom) == (55, 34)
    out = str(tmp_path / "l.pdf")
    cust = dict(db.customer("90000003"), has_pn=1)
    stages.render_letter(db, folder, "1004PN", cust, 1000, "S-1", "2026-10-20", None, out)
    from pypdf import PdfReader
    assert len(PdfReader(out).pages) == 1          # the legal letter stays on one page of the letter paper
